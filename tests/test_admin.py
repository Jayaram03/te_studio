import csv
import io
from datetime import date, timedelta

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app import api, company, db, pipeline, trips
from app.config import get_settings
from conftest import signed_in_client, sample

KASHMIR = "Off_Season_2026_Kashmir_package_5N_6D.pdf"


def _setup(s, ex):
    doc, _ = pipeline.ingest(s, *sample(KASHMIR), ex)
    pipeline.approve(s, doc, {"rate_type": "net"})
    pkg = s.scalar(select(db.Package))
    t = trips.create_trip(s, {"title": "Kashmir for the Iyers", "customer_name": "Mrs Iyer", "customer_phone": "+91 90000 00000",
                              "start_date": "2026-09-10", "adults": 2, "markup_pct": 10, "gst_pct": 5,
                              "lead_source": "Instagram", "assigned_to": "Rakesh"})
    trips.apply_package(s, t, pkg.id, "Deluxe")
    return t, pkg


def test_pipeline_followups_payments_and_dashboard(session, extractor):
    t, _ = _setup(session, extractor)
    today = date(2026, 9, 1)
    trips.update_trip(session, t, {"status": "quoted", "follow_up_date": (today - timedelta(days=2)).isoformat()})
    assert "Status changed: enquiry -> quoted" in t.notes_log[-1].text
    trips.add_note(session, t, "Called, wants houseboat upgrade", "Rakesh")
    d = trips.dashboard(session, today)
    assert d["follow_ups"][0]["id"] == t.id and d["follow_ups"][0]["overdue_days"] == 2
    assert d["pipeline"]["quoted"]["count"] == 1 and d["pipeline"]["quoted"]["value"] == round(34000 * 1.1 * 1.05, 2)

    trips.update_trip(session, t, {"status": "confirmed", "follow_up_date": ""})
    trips.add_payment(session, t, 20000, "2026-09-02", "UPI")
    v = trips.trip_view(session, t)
    assert v["totals"]["paid"] == 20000 and v["totals"]["balance"] == round(34000 * 1.1 * 1.05 - 20000, 2)
    assert v["follow_up_date"] is None
    d = trips.dashboard(session, today)
    assert not d["follow_ups"] and d["upcoming_departures"][0]["id"] == t.id
    assert d["lead_sources"][0] == {"source": "Instagram", "trips": 1, "confirmed": 1, "value": round(34000 * 1.1 * 1.05, 2)}
    assert d["conversion_pct"] == 100.0 and d["balance_due"] == v["totals"]["balance"]


def test_duplicate_keeps_plan_not_customer(session, extractor):
    t, _ = _setup(session, extractor)
    trips.add_payment(session, t, 5000)
    c = trips.duplicate(session, t)
    v = trips.trip_view(session, c)
    assert v["title"] == "Copy of Kashmir for the Iyers" and v["customer_name"] is None and v["status"] == "enquiry"
    assert len(v["days"]) == 6 and v["totals"]["cost"] == 34000 and v["totals"]["paid"] == 0


def test_pdf_csv_xlsx_exports(session, extractor):
    t, pkg = _setup(session, extractor)
    api.set_extractor(extractor)
    c = signed_in_client(session)
    r = c.get(f"/api/trips/{t.id}/pdf")
    assert r.status_code == 200 and r.content[:4] == b"%PDF" and len(r.content) > 20000
    assert "Quote-" in r.headers["content-disposition"]
    assert c.get(f"/api/trips/{t.id}/pdf", params={"prices": False}).content[:4] == b"%PDF"
    assert c.get(f"/api/catalog/packages/{pkg.id}/pdf").content[:4] == b"%PDF"
    rows = list(csv.reader(io.StringIO(c.get(f"/api/trips/{t.id}/csv").content.decode("utf-8-sig"))))
    assert rows[0][0] == "Type" and any(r[2] == "Selling price" for r in rows)
    itin = list(csv.reader(io.StringIO(c.get(f"/api/trips/{t.id}/csv", params={"part": "itinerary"}).content.decode("utf-8-sig"))))
    assert len(itin) == 7 and itin[1][2] == "Arrival Srinagar"
    prices = list(csv.reader(io.StringIO(c.get(f"/api/catalog/packages/{pkg.id}/csv").content.decode("utf-8-sig"))))
    assert len(prices) == 37
    for what in ("hotel-rates", "hotels", "packages", "services", "places", "suppliers", "trips"):
        assert c.get(f"/api/export/{what}.csv").status_code == 200
    wb = load_workbook(io.BytesIO(c.get("/api/export/library.xlsx").content))
    assert wb.sheetnames == ["Hotel rates", "Packages", "Activities & transfers", "Hotels", "Places", "Suppliers", "Trips"]
    assert wb["Hotels"].max_row == 111
    wb = load_workbook(io.BytesIO(c.get(f"/api/trips/{t.id}/xlsx").content))
    assert wb.sheetnames == ["Itinerary", "Costing"]


def test_share_link_is_public_and_revocable(session, extractor):
    t, _ = _setup(session, extractor)
    c = signed_in_client(session)
    v = c.post(f"/api/trips/{t.id}/share", json={"enable": True, "prices": True}).json()
    tok = v["share_token"]
    anon = TestClient(api.app)                                # a customer: not signed in
    page = anon.get(f"/share/{tok}")
    assert page.status_code == 200 and "Kashmir for the Iyers" in page.text and "₹" in page.text
    assert "Rakesh" in page.text and "+91 89397 18676" in page.text   # preparer from Settings team
    assert anon.get(f"/share/{tok}/pdf").content[:4] == b"%PDF"
    assert anon.get("/api/trips").status_code == 401                 # but nothing else
    c.post(f"/api/trips/{t.id}/share", json={"prices": False})
    assert "Per person" not in anon.get(f"/share/{tok}").text
    c.post(f"/api/trips/{t.id}/share", json={"enable": False})
    assert anon.get(f"/share/{tok}").status_code == 404


def test_settings_roundtrip(session):
    s = company.get_all(session)
    assert s["company"]["email"] == "travelepisodeschennai@gmail.com" and len(s["team"]) == 3
    company.save(session, {"quote": {**s["quote"], "default_markup_pct": 12, "bank_details": "HDFC 123"}})
    s2 = company.get_all(session)
    assert s2["quote"]["default_markup_pct"] == 12 and s2["quote"]["bank_details"] == "HDFC 123"
    c = signed_in_client(session)
    t = c.post("/api/trips", json={"title": "x"}).json()
    assert t["markup_pct"] == 12 and t["status"] == "enquiry"
    row = c.get("/api/trips").json()[0]
    assert {"sell", "balance", "assigned_to", "follow_up_date", "lead_source"} <= set(row)
