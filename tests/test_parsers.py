"""Every supported format is read into text + tables, whatever its layout; bad files give clear messages."""
import io

import pytest

from app import pipeline
from app.parsers import MAX_ROWS_PER_PART, UnsupportedFile, parse_file, sniff


def _docx(with_table=True, header="Valid 1 Oct 2026 - 31 Mar 2027"):
    import docx
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = header
    d.add_heading("Green Valley Holidays - Kerala rates", 1)
    d.add_paragraph("All rates are net B2B in INR, taxes extra.")
    if with_table:
        t = d.add_table(rows=3, cols=3)
        t.cell(0, 0).merge(t.cell(0, 1)).text = "Peak season"       # merged header over two columns
        t.cell(0, 2).text = "Regular"
        for c, v in enumerate(["Deluxe CP", "7500", "5000"]):
            t.cell(1, c).text = v
        for c, v in enumerate(["Suite CP", "9900", "7200"]):
            t.cell(2, c).text = v
    d.add_paragraph("Child below 5 free.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_word_document_keeps_order_tables_merged_cells_and_header():
    doc = parse_file("rates.docx", _docx())
    assert doc.kind == "word"
    t = doc.text
    assert "Valid 1 Oct 2026 - 31 Mar 2027" in t                     # page header kept
    assert "| Peak season | Peak season | Regular |" in t              # merged cell repeated per column
    assert "| Deluxe CP | 7500 | 5000 |" in t
    assert t.index("net B2B") < t.index("Deluxe CP") < t.index("Child below 5")   # document order


def test_csv_with_semicolons_and_windows_encoding():
    data = "Hotel;Room;Meal;Rate\nCaf\xe9 Munnar;Deluxe;CP;4500\n".encode("cp1252")
    doc = parse_file("export.csv", data)
    assert doc.meta["delimiter"] == ";" and "| Café Munnar | Deluxe | CP | 4500 |" in doc.text


def test_long_sheet_is_split_with_the_header_repeated():
    rows = ["Hotel,Room,Rate"] + [f"H{i},Deluxe,{1000 + i}" for i in range(MAX_ROWS_PER_PART * 2)]
    doc = parse_file("big.csv", "\n".join(rows).encode())
    assert len(doc.chunks) == 3
    assert all("| Hotel | Room | Rate |" in c for c in doc.chunks) and "part 2 of 3" in doc.chunks[1]


def test_xml_is_read_safely():
    xml = b"""<?xml version="1.0"?><rates supplier="Siam Link"><hotel name="Patong Beach"><room type="Deluxe">
    <season from="2026-11-01" to="2027-04-30">3200</season></room></hotel></rates>"""
    t = parse_file("feed.xml", xml).text
    assert 'rates [supplier="Siam Link"]' in t and 'season [from="2026-11-01" to="2027-04-30"]: 3200' in t
    bomb = b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa"><!ENTITY b "&a;&a;&a;">]><x>&b;</x>'
    with pytest.raises(UnsupportedFile):
        parse_file("bomb.xml", bomb)                                     # entity expansion refused


def test_html_export_named_xls_is_read_as_html():
    html = (b"<html><head><style>x{}</style><script>alert(1)</script></head><body><h1>Rates</h1>"
            b"<table><tr><th colspan=2>Peak</th></tr><tr><td>Deluxe</td><td>7500</td></tr></table></body></html>")
    assert sniff(html, ".xls") == ".html"
    doc = parse_file("booking_system_export.xls", html)
    assert doc.kind == "html" and "| Peak | Peak |" in doc.text and "| Deluxe | 7500 |" in doc.text
    assert "alert" not in doc.text


def test_content_beats_the_extension():
    assert sniff(_docx(), ".pdf") == ".docx"
    assert sniff(b"\x89PNG\r\n\x1a\nxxxx", ".pdf") == ".png"
    assert parse_file("scan.pdf", b"\x89PNG\r\n\x1a\nxxxx").kind == "image"


def test_json_and_text():
    assert "Deluxe" in parse_file("rates.json", b'{"rooms": [{"name": "Deluxe", "rate": 4500}]}').text
    assert parse_file("notes.txt", "Houseboat 1 night: 9500".encode()).kind == "text"


@pytest.mark.parametrize("name,data,msg", [
    ("old.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 100, "Save as"),
    ("empty.pdf", b"", "empty"),
    ("broken.xlsx", b"PK\x03\x04 not really a zip", "Could not read"),
    ("mystery.exe", b"MZ\x90\x00", "Unsupported"),
])
def test_unreadable_files_get_clear_messages(name, data, msg):
    with pytest.raises((UnsupportedFile, ValueError)) as e:
        parse_file(name, data)
    assert msg.lower() in str(e.value).lower()


def test_upload_refuses_unsupported_and_old_word_files_before_storing(session):
    with pytest.raises(pipeline.PipelineError, match="Save as"):
        pipeline.intake(session, "old.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 100)
    with pytest.raises(pipeline.PipelineError, match="Unsupported"):
        pipeline.intake(session, "x.exe", b"MZ\x90\x00")


def test_word_file_reaches_the_ai_as_text(session):
    seen = {}

    class Spy:
        model_name = "spy"

        def extract(self, doc, filename, supplier_hint=None, reference=None):
            seen["text"] = doc.text
            return {"document_type": "hotel_rate_sheet"}
    doc, _ = pipeline.ingest(session, "Green Valley rates.docx", _docx(), Spy())
    assert doc.status == "needs_review" and "| Deluxe CP | 7500 | 5000 |" in seen["text"]
