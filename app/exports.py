"""CSV and Excel exports of the library and of trips."""
from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import db, trips


def _f(x):
    return None if x is None else float(x)


def rows_hotel_rates(s: Session):
    q = (select(db.HotelRate, db.Hotel, db.RoomType).join(db.Hotel, db.HotelRate.hotel_id == db.Hotel.id)
         .join(db.RoomType, db.HotelRate.room_type_id == db.RoomType.id).where(db.HotelRate.status == "active")
         .order_by(db.Hotel.city, db.Hotel.name, db.RoomType.name, db.HotelRate.valid_from))
    yield ["Hotel", "City", "Room", "Meal plan", "Occupancy", "Basis", "Amount", "Currency", "Net (B2B)",
           "Taxes included", "Valid from", "Valid to", "Days", "Season", "Min nights", "Source document"]
    for r, h, rt in s.execute(q).all():
        yield [h.name, h.city, rt.name, r.meal_plan, r.occupancy, r.basis, _f(r.amount), r.currency, r.is_net,
               r.taxes_included, r.valid_from, r.valid_to, r.weekdays or "all", r.season_name, r.min_nights,
               r.source_document_id]


def rows_hotels(s: Session):
    yield ["Hotel", "City", "Destination", "Category", "Type", "Stars", "Address"]
    for h in s.scalars(select(db.Hotel).order_by(db.Hotel.city, db.Hotel.name)).all():
        yield [h.name, h.city, h.destination, h.category, h.property_type, h.star_rating, h.address]


def rows_packages(s: Session, package_id: int | None = None):
    yield ["Package", "Supplier", "Region", "Nights", "Days", "Category", "Pax from", "Pax to", "Occupancy",
           "Basis", "Amount", "Currency", "Net (B2B)", "Valid from", "Valid to", "Season"]
    q = select(db.Package).where(db.Package.status == "active")
    if package_id:
        q = select(db.Package).where(db.Package.id == package_id)
    for p in s.scalars(q.order_by(db.Package.title)).all():
        for x in sorted(p.prices, key=lambda x: (x.category or "", x.occupancy != "double", x.pax_min or 0)):
            yield [p.title, p.supplier.name if p.supplier else None, p.region, p.nights, p.days, x.category,
                   x.pax_min, x.pax_max, x.occupancy, x.basis, _f(x.amount), x.currency, x.is_net, x.valid_from,
                   x.valid_to, x.season_name]


def rows_package_itinerary(s: Session, package_id: int):
    p = s.get(db.Package, package_id)
    yield ["Day", "Title", "Overnight", "Meals", "Description"]
    for d in p.days_list:
        yield [d.day_number, d.title, d.overnight, d.meals, d.description]


def rows_services(s: Session):
    yield ["Kind", "Service", "Destination", "Vehicle", "Basis", "Seats", "Amount", "Up to", "Currency", "Add-on",
           "Net (B2B)", "Valid from", "Valid to", "Notes"]
    for r in s.scalars(select(db.ServiceRate).where(db.ServiceRate.status == "active")
                       .order_by(db.ServiceRate.destination, db.ServiceRate.name)).all():
        yield [r.kind, r.name, r.destination, r.vehicle_type, r.basis, r.pax_max, _f(r.amount), _f(r.amount_max),
               r.currency, r.optional, r.is_net, r.valid_from, r.valid_to, r.notes]


def rows_places(s: Session):
    yield ["Place", "Destination", "Region", "Kind", "Availability", "Description"]
    for p in s.scalars(select(db.Place).order_by(db.Place.destination, db.Place.name)).all():
        yield [p.name, p.destination, p.region, p.kind, p.availability, p.description]


def rows_suppliers(s: Session):
    yield ["Supplier", "Type", "City", "Country", "Contact", "Phone", "Email", "Address", "Website", "GSTIN"]
    for x in s.scalars(select(db.Supplier).order_by(db.Supplier.name)).all():
        yield [x.name, x.type, x.city, x.country, x.contact_person, x.phone, x.email, x.address, x.website, x.gst_number]


def rows_trips(s: Session):
    yield ["Trip #", "Title", "Customer", "Phone", "Email", "Lead source", "Assigned to", "Status", "Destination",
           "Start date", "Adults", "Children", "Cost", "Markup", "GST", "Selling price", "Paid", "Balance",
           "Currency", "Follow-up date", "Created"]
    for t in s.scalars(select(db.Trip).order_by(db.Trip.id)).all():
        tt = trips.totals(t)
        yield [t.id, t.title, t.customer_name, t.customer_phone, t.customer_email, t.lead_source, t.assigned_to,
               t.status, t.destination, t.start_date, t.adults, t.children_with_bed + t.children_without_bed,
               tt["cost"], tt["markup"], tt["gst"], tt["sell"], tt["paid"], tt["balance"], tt["currency"],
               t.follow_up_date, t.created_at and t.created_at.date()]


def rows_trip_costing(s: Session, t: db.Trip):
    v = trips.trip_view(s, t)
    yield ["Type", "Day", "Item", "Quantity", "Unit price", "Amount", "Currency", "Add-on (not in total)"]
    for i in v["items"]:
        yield [i["kind"], i["day_position"], i["description"], i["quantity"], i["unit_amount"], i["amount"],
               i["currency"], i["optional"]]
    T = v["totals"]
    for label, key in (("Cost", "cost"), (f"Markup {v['markup_pct']:g}%", "markup"), (f"GST {v['gst_pct']:g}%", "gst"),
                       ("Selling price", "sell"), ("Per person", "per_person"), ("Paid", "paid"), ("Balance", "balance")):
        yield ["total", None, label, None, None, T[key], T["currency"], None]


def rows_trip_itinerary(s: Session, t: db.Trip):
    v = trips.trip_view(s, t)
    yield ["Day", "Date", "Title", "Overnight", "Meals", "Description"]
    for d in v["days"]:
        yield [d["position"], d["date"], d["title"], d["overnight"], d["meals"], d["description"]]


EXPORTS = {"hotel-rates": rows_hotel_rates, "hotels": rows_hotels, "packages": rows_packages,
           "services": rows_services, "places": rows_places, "suppliers": rows_suppliers, "trips": rows_trips}


def to_csv(rows) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    for r in rows:
        w.writerow(["" if v is None else ("yes" if v is True else "no" if v is False else v) for v in r])
    return ("﻿" + buf.getvalue()).encode("utf-8")      # BOM so Excel opens Indian text correctly


def to_xlsx(sheets: dict[str, object]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="E2E3F3")
    for title, rows in sheets.items():
        ws = wb.create_sheet(title[:31])
        for i, r in enumerate(rows):
            ws.append([("yes" if v is True else "no" if v is False else v) for v in r])
            if i == 0:
                for c in ws[1]:
                    c.font = Font(bold=True, color="4F4C4D")
                    c.fill = head_fill
        ws.freeze_panes = "A2"
        for col in ws.columns:
            width = min(max(len(str(c.value or "")) for c in col[:200]) + 2, 60)
            ws.column_dimensions[get_column_letter(col[0].column)].width = max(width, 8)
        ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def library_xlsx(s: Session) -> bytes:
    return to_xlsx({"Hotel rates": rows_hotel_rates(s), "Packages": rows_packages(s), "Activities & transfers":
                    rows_services(s), "Hotels": rows_hotels(s), "Places": rows_places(s),
                    "Suppliers": rows_suppliers(s), "Trips": rows_trips(s)})
